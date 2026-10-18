"""Check the engine against Ayres and Nalebuff's own simulation workbook.

The authors posted the spreadsheet behind their 2010 book at
lifecycleinvesting.net (Core_simulation.xlsx, 1871-2009, ages 23-66, an 83%
target, call money as the borrowing rate). It carries no license, so it is
downloaded into a local ``.cache/`` and never committed.

Two runs on the workbook's own monthly real returns and contributions:

1. The workbook's timing, copied from its formulas: next month's contribution
   is added before this month's return is applied, and the allocation divides
   by savings plus the current contribution. This reproduces the workbook's
   Summary sheet exactly.
2. ``lifecycle_leverage.simulate``, which uses standard timing (a month's
   contribution earns that month's return). It lands within about 0.6%,
   the same small offset for every strategy.

Usage:  python validate_workbook.py      (needs openpyxl)
"""

from __future__ import annotations

import io
import urllib.request
import zipfile
from pathlib import Path

import numpy as np
import openpyxl

from lifecycle_leverage import (
    Constant,
    Glide,
    Target,
    pv_remaining,
    simulate,
    summarize,
)

HERE = Path(__file__).resolve().parent
CACHE = HERE / ".cache"
URL = "http://lifecycleinvesting.net/Data/Core_simulation.zip"
WORKBOOK_SUMMARY = {
    "lifecycle": {"median": 1146812, "mean": 1223105, "std": 453379, "min": 387172,
                  "p10": 701834, "p90": 1929577, "ce2": 1044919, "ce4": 870326},
    "constant75": {"median": 691427, "mean": 748839, "min": 308726},
    "birthday9050": {"median": 641555, "mean": 646575, "min": 290310},
}  # fmt: skip
LAMBDA = 0.83030407
T = 528


def load():
    CACHE.mkdir(exist_ok=True)
    xlsx = CACHE / "Core_simulation.xlsx"
    if not xlsx.exists():
        raw = urllib.request.urlopen(URL, timeout=600).read()  # noqa: S310
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            xlsx.write_bytes(z.read("Core_simulation.xlsx"))
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    rows = list(wb["Monthly data"].iter_rows(min_row=2, max_row=1669, values_only=True))
    bond = np.array([r[11] for r in rows], dtype=float)
    margin = np.array([r[12] for r in rows], dtype=float)
    stock = np.array([r[13] for r in rows], dtype=float)
    contrib = np.array(
        [
            r[6]
            for r in wb["Contributions"].iter_rows(
                min_row=2, max_row=T + 1, values_only=True
            )
        ],
        dtype=float,
    )
    s = wb["Summary"]
    r_lev = (1 + s["B4"].value) ** (1 / 12) - 1
    r_unlev = (1 + s["B5"].value) ** (1 / 12) - 1
    return stock, bond, margin, contrib, r_lev, r_unlev


def workbook_timing(stock, bond, margin, contrib, r_lev, r_unlev):
    pv_lev = pv_remaining(contrib, r_lev)
    pv_unlev = pv_remaining(contrib, r_unlev)
    c = np.append(contrib, 0.0)
    finals = []
    for k in range(96):
        s0, saved, prev = 12 * k, 0.0, 0.0
        for t in range(T):
            saved = (saved + c[t]) * (1 + prev)
            x = saved + c[t]
            want = LAMBDA * (pv_lev[t] + saved) / x
            lam = (
                min(2.0, want)
                if want > 1
                else min(1.0, LAMBDA * (pv_unlev[t] + saved) / x)
            )
            m = s0 + t
            prev = max(
                -1.0, lam * stock[m] + (1 - lam) * (margin[m] if lam > 1 else bond[m])
            )
        finals.append((saved + c[T]) * (1 + prev))
    return np.array(finals)


def main() -> None:
    stock, bond, margin, contrib, r_lev, r_unlev = load()
    exact = summarize(workbook_timing(stock, bond, margin, contrib, r_lev, r_unlev))
    print("1) Workbook timing vs the workbook's Summary sheet (lifecycle strategy)")
    for k, v in WORKBOOK_SUMMARY["lifecycle"].items():
        print(f"   {k:6s} ours {exact[k]:>12,.0f}   workbook {v:>12,.0f}")
    idx = np.array([np.arange(12 * k, 12 * k + T) for k in range(96)])
    print("2) lifecycle_leverage.simulate (standard timing)")
    pols = {"lifecycle": Target(LAMBDA, cap=2.0), "constant75": Constant(0.75),
            "birthday9050": Glide(0.9, 0.5)}  # fmt: skip
    for name, pol in pols.items():
        w = simulate(
            stock[idx],
            bond[idx],
            margin[idx],
            contrib,
            pol,
            r_lev=r_lev,
            r_unlev=r_unlev,
        ).wealth
        s = summarize(w)
        ref = WORKBOOK_SUMMARY[name]
        print(
            f"   {name:12s} " + "  ".join(f"{k} {s[k] / ref[k] - 1:+.2%}" for k in ref)
        )


if __name__ == "__main__":
    main()
